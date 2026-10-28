"""DOCX parsing with python-docx.

DOCX has no fixed pages, so page numbers are derived as follows (recorded as `page_mode`):

- "rendered": Word stores `w:lastRenderedPageBreak` markers where its last layout broke pages.
  When present we split on them (plus explicit page breaks), which matches the printed page
  numbers closely enough for citations.
- "virtual": files not saved by Word (e.g. generated) have no layout markers. We split on
  explicit page breaks and otherwise every ~N characters at paragraph boundaries, preferring
  to start a new page at a heading.

Converting to PDF with LibreOffice would give exact pages but adds a heavy dependency; that
can be revisited if page accuracy in the Phase 9 evaluation turns out to be a problem.
"""

import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.parsers.base import Block, Page, PageMode, ParsedDocument, ParseError

BreakPosition = Literal["before", "after"]

_HEADING_STYLE = re.compile(r"heading (\d)", re.IGNORECASE)
_W_T = qn("w:t")
_W_BR = qn("w:br")
_W_TYPE = qn("w:type")
_W_RENDERED_BREAK = qn("w:lastRenderedPageBreak")
_W_PAGE_BREAK_BEFORE = qn("w:pageBreakBefore")
_W_OUTLINE_LVL = qn("w:outlineLvl")
_W_VAL = qn("w:val")


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name == "Title":
        return 1
    if m := _HEADING_STYLE.fullmatch(style_name or ""):
        return int(m.group(1))
    p_pr = paragraph._p.pPr
    if p_pr is not None and (outline := p_pr.find(_W_OUTLINE_LVL)) is not None:
        level = int(outline.get(_W_VAL, "9"))
        if level < 9:  # 9 means "body text"
            return level + 1
    return None


def _page_break(paragraph: Paragraph) -> tuple[BreakPosition | None, bool]:
    """Where a page break falls relative to this paragraph, and whether it was Word-rendered."""
    seen_text = False
    for el in paragraph._p.iter():
        if el.tag == _W_T and el.text:
            seen_text = True
        elif el.tag == _W_RENDERED_BREAK or (el.tag == _W_BR and el.get(_W_TYPE) == "page"):
            return ("after" if seen_text else "before"), el.tag == _W_RENDERED_BREAK
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.find(_W_PAGE_BREAK_BEFORE) is not None:
        return "before", False
    return None, False


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = cell.text.strip()
            if not cells or cells[-1] != text:  # merged cells repeat their text
                cells.append(text)
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _paginate(
    items: list[tuple[Block, BreakPosition | None]], *, by_breaks_only: bool, page_chars: int
) -> list[list[Block]]:
    pages: list[list[Block]] = [[]]
    size = 0
    for block, brk in items:
        overflow = (
            not by_breaks_only
            and pages[-1]
            and (
                size + len(block.text) > page_chars
                or (block.is_heading and size >= page_chars // 2)
            )
        )
        if (brk == "before" or overflow) and pages[-1]:
            pages.append([])
            size = 0
        pages[-1].append(block)
        size += len(block.text)
        if brk == "after":
            pages.append([])
            size = 0
    if len(pages) > 1 and not pages[-1]:
        pages.pop()
    return pages


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value if value.tzinfo else value.replace(tzinfo=UTC)  # DOCX stores UTC


def parse_docx(path: Path, *, virtual_page_chars: int) -> ParsedDocument:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError("The DOCX file could not be opened; the file may be corrupt.") from exc

    items: list[tuple[Block, BreakPosition | None]] = []
    rendered = False
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            level = _heading_level(item)
            brk, is_rendered = _page_break(item)
            rendered = rendered or is_rendered
            items.append((Block(item.text, heading_level=level, reflow=level is not None), brk))
        elif isinstance(item, Table):
            items.append((Block(_table_text(item)), None))

    page_mode: PageMode = "rendered" if rendered else "virtual"
    pages = _paginate(items, by_breaks_only=rendered, page_chars=virtual_page_chars)

    props = document.core_properties
    extra: dict[str, Any] = {
        "subject": props.subject or None,
        "keywords": props.keywords or None,
        "last_modified_by": props.last_modified_by or None,
        "modified_at": (d.isoformat() if (d := _as_utc(props.modified)) else None),
        "table_count": len(document.tables),
    }
    return ParsedDocument(
        pages=[Page(i + 1, blocks) for i, blocks in enumerate(pages)],
        page_mode=page_mode,
        title=(props.title or "").strip() or None,
        author=(props.author or "").strip() or None,
        created_at=_as_utc(props.created),
        extra={k: v for k, v in extra.items() if v is not None},
    )
