"""Builders for test documents, generated on the fly so no binary files are committed."""

from pathlib import Path

import docx
import pymupdf

BODY = (
    "Employees may work remotely up to three days per week with manager approval. "
    "Equipment is provided by the IT department and remains company property."
)


def handbook_pdf(path: Path, pages: int = 4) -> Path:
    """A multi-page PDF with a running header, a page-number footer, headings and body text."""
    doc = pymupdf.open()
    for n in range(1, pages + 1):
        page = doc.new_page()
        page.insert_text((72, 40), "ACME Corp Employee Handbook", fontsize=9)
        page.insert_text((72, 100), f"Section {n}: Policy Area {n}", fontsize=18)
        page.insert_textbox(pymupdf.Rect(72, 120, 520, 300), f"{BODY} (page {n})", fontsize=11)
        page.insert_text((280, 800), f"Page {n} of {pages}", fontsize=9)
    doc.set_metadata(
        {
            "title": "ACME Employee Handbook",
            "author": "HR Team",
            "creationDate": "D:20240115093000+01'00'",
            "subject": "Policies",
        }
    )
    doc.save(path)
    return path


def two_column_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()  # 595 x 842
    page.insert_textbox(pymupdf.Rect(72, 50, 523, 90), "Full Width Title", fontsize=20)
    left = ["Left one starts the article.", "Left two continues the story."]
    right = ["Right one comes after the left column.", "Right two ends the page."]
    for i, (lt, rt) in enumerate(zip(left, right, strict=True)):
        y = 120 + i * 80
        page.insert_textbox(pymupdf.Rect(50, y, 280, y + 60), lt, fontsize=11)
        page.insert_textbox(pymupdf.Rect(315, y, 545, y + 60), rt, fontsize=11)
    doc.save(path)
    return path


def hyphenated_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(
        pymupdf.Rect(72, 72, 520, 200),
        "This is an exam-\nple of a hyphen-\nated line and the COVID-\n19 rules.",
        fontsize=11,
    )
    doc.save(path)
    return path


def outline_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    for title in ("Introduction", "Scope"):
        page = doc.new_page()
        page.insert_text((72, 80), title, fontsize=11)  # same size as body: only the outline knows
        page.insert_textbox(pymupdf.Rect(72, 100, 520, 300), BODY, fontsize=11)
    doc.set_toc([[1, "Introduction", 1], [2, "Scope", 2]])
    doc.save(path)
    return path


def empty_pdf(path: Path, pages: int = 2) -> Path:
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page()
    doc.save(path)
    return path


def sample_docx(path: Path, *, page_break: bool = True) -> Path:
    document = docx.Document()
    document.core_properties.title = "Travel Policy"
    document.core_properties.author = "Finance"
    document.add_heading("Travel Policy", level=1)
    document.add_paragraph("All travel must be booked through the approved agency.")
    document.add_heading("Expenses", level=2)
    document.add_paragraph("Meals are reimbursed up to 50 EUR per day.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "Hotel", "150 EUR"
    if page_break:
        document.add_page_break()
    document.add_heading("Approvals", level=2)
    document.add_paragraph("Trips abroad need director approval.")
    document.save(str(path))
    return path


def long_docx(path: Path, paragraphs: int = 40) -> Path:
    document = docx.Document()
    for i in range(paragraphs):
        document.add_paragraph(f"Paragraph {i}. " + BODY)
    document.save(str(path))
    return path
